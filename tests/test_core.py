import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

import numpy as np

from rag.config import Config, load_env
from rag.evidence import select_output
from rag.models import decode_json, verify_snapshot
from rag.parsers import chunk_text, parse_corpus
from rag.pipeline import Pipeline
from rag.protocol import request
from rag.retrieval import Retriever, rrf
from rag.schema import Evidence, validate_output

ROOT = Path(__file__).resolve().parents[1]


def config_at(temp):
    return Config(backend="extractive", device="cpu", require_rocm=False,
                  state_dir=str(temp/"state"), output_dir=str(temp/"output"),
                  model_dir=str(temp/"models"), socket=str(temp/"worker.sock"))


def internal(answer, proofs, links=None):
    return {"answer": answer, "evidence_ids": [p[0] for p in proofs],
            "proofs": [{"evidence_id": i, "quote": q, "role": role} for i,q,role in proofs],
            "links": links or []}


class EvidenceTests(unittest.TestCase):
    def test_chain_preserves_anchor_and_value(self):
        a = Evidence("a", "events/z.log", "PX-713 thermal throttle: BUG-6723", ".log")
        b = Evidence("b", "fixes/renamed.csv", "BUG-6723 fixed firmware: 8.2.5", ".csv")
        value = internal("8.2.5", [("a", a.content, "anchor"), ("b", b.content, "value")],
                         [{"from_id":"a", "to_id":"b", "identifier":"BUG-6723"}])
        output, reason = select_output(value, [a,b], "Which firmware fixed the PX-713 thermal throttle?")
        self.assertEqual(output, {"answer":"8.2.5", "citations":["events/z.log","fixes/renamed.csv"]})
        self.assertEqual(reason, "answered")
        value["links"] = []
        self.assertEqual(select_output(value,[a,b],"PX-713 throttle")[0]["answer"], "")

    def test_rejects_hallucinated_paths_quotes_and_extra_sources(self):
        a = Evidence("a", "a.txt", "PX-713 revision: REV-D4", ".txt")
        b = Evidence("b", "b.txt", "Unrelated board: REV-D4", ".txt")
        for value in [internal("REV-D4", [("unknown",a.content,"anchor")]),
                      internal("REV-D4", [("a","PX-713 revision: REV-D9","anchor")]),
                      internal("REV-D4", [("a",a.content,"anchor"),("b",b.content,"value")])]:
            self.assertEqual(select_output(value,[a,b],"PX-713 revision")[0]["answer"], "")

    def test_qualifiers_must_be_complete(self):
        for text, full, partial in [("PX-713 revision: REV-D4","REV-D4","D4"),
                                     ("PX-713 quarter: Q2 FY34","Q2 FY34","Q2")]:
            a = Evidence("a", "a.txt", text, ".txt")
            self.assertTrue(select_output(internal(full,[("a",text,"anchor")]),[a],"PX-713")[0]["answer"])
            self.assertFalse(select_output(internal(partial,[("a",text,"anchor")]),[a],"PX-713")[0]["answer"])

    def test_schema(self):
        for output in [{"answer":1,"citations":[]},{"answer":"x","citations":"a"},
                       {"answer":"","citations":["a"]},{"answer":"x","citations":[]}]:
            with self.assertRaises(ValueError): validate_output(output)
        self.assertEqual(decode_json('```json\n{"answer":""}\n```'), {"answer":""})

    def test_truncated_numbers_and_link_identifiers_rejected(self):
        a = Evidence("a","a.txt","PX-713 timeout: 235", ".txt")
        self.assertEqual(select_output(internal("23",[("a",a.content,"anchor")]),[a],"PX-713 timeout")[0]["answer"],"")
        a = Evidence("a","a.txt","PX-713 throttle BUG-6723", ".txt")
        b = Evidence("b","b.txt","BUG-6799 fixed firmware: 8.2.5", ".txt")
        value = internal("8.2.5",[("a",a.content,"anchor"),("b",b.content,"value")],
                         [{"from_id":"a","to_id":"b","identifier":"BUG-67"}])
        self.assertEqual(select_output(value,[a,b],"PX-713 throttle")[0]["answer"],"")


class RetrievalTests(unittest.TestCase):
    def test_rrf_and_hybrid_are_deterministic(self):
        self.assertEqual(rrf([[2,1],[1,2]],60), [1,2])
        self.assertEqual(rrf([[7,2],[2,3]],60)[0], 2)
        evidence = [Evidence("a","x.txt","PX-713 ticket BUG-6723", ".txt"),
                    Evidence("b","y.txt","BUG-6723 fixed firmware 8.2.5", ".txt"),
                    Evidence("c","noise.txt","fan assembly", ".txt")]
        class Encoder:
            def embed(self, texts, query=False):
                return np.array([[1.,0.] for _ in texts])
            def rerank(self, query, texts):
                return [float("BUG-6723" in t) for t in texts]
        for method in ["bm25","dense","hybrid"]:
            c = Config(method=method, context_k=2)
            r = Retriever(evidence,c,Encoder(),np.array([[1.,0.],[.8,.6],[0.,1.]]))
            found,_ = r.search("PX-713",time.monotonic()+5)
            self.assertEqual(found[0].chunk_id, "a")
        c = Config(method="bm25",expand=True,context_k=3)
        found,trace = Retriever(evidence,c,Encoder()).search("PX-713 thermal throttle firmware",time.monotonic()+5)
        self.assertIn("BUG-6723",trace["expansion_identifiers"])
        self.assertIn("y.txt",[e.source_path for e in found])

    def test_long_line_chunking(self):
        blocks = list(chunk_text(" ".join(f"word{i}" for i in range(500)),100,20))
        self.assertGreater(len(blocks),1)
        self.assertLessEqual(max(len(b[0].split()) for b in blocks),100)


class ParserTests(unittest.TestCase):
    def test_mixed_formats_failures_and_provenance(self):
        from docx import Document
        from openpyxl import Workbook
        import pymupdf
        from PIL import Image
        with tempfile.TemporaryDirectory() as temp:
            temp = Path(temp)
            corpus = temp/"new"/"arbitrary"
            corpus.mkdir(parents=True)
            (corpus/"empty").mkdir()
            (corpus/"safe.py").write_text("TIMEOUT = 235\nraise RuntimeError('do not execute')")
            (corpus/"table.csv").write_text("Component;Part number\nZ;FAN-8741\n")
            (corpus/"unknown.dat").write_bytes(b"unknown")
            (corpus/"broken.docx").write_bytes(b"invalid")
            denied = corpus/"denied.txt"
            denied.write_text("unreadable secret")
            denied.chmod(0)
            doc = Document()
            doc.add_paragraph("Roadmap")
            t = doc.add_table(rows=1,cols=2)
            t.rows[0].cells[0].text,t.rows[0].cells[1].text = "Device","Sampling"
            t.add_row().cells[0].text = "Z"
            t.rows[1].cells[1].text = "Q2 FY34"
            doc.save(corpus/"road.docx")
            book = Workbook()
            book.active.append(["Name","Value"])
            book.active.append(["First","A"])
            sh = book.create_sheet("Second")
            sh.append(["Name","Value"])
            sh.append(["Fan","FAN-8741"])
            book.save(corpus/"book.xlsx")
            pdf = pymupdf.open()
            pdf.new_page().insert_text((30,50), "Maximum temperature: 87 C")
            pdf.save(corpus/"spec.pdf")
            pdf.close()
            pdf = pymupdf.open()
            pdf.new_page().insert_text((30,50), "never retrieve this")
            pdf.save(corpus/"locked.pdf", encryption=pymupdf.PDF_ENCRYPT_AES_256,owner_pw="owner",user_pw="password")
            pdf.close()
            pdf = pymupdf.open()
            pdf.new_page()
            pdf.save(corpus/"scan.pdf")
            pdf.close()
            Image.new("RGB",(100,100),"white").save(corpus/"label.png")
            calls = []
            def fake_vision(path, deadline):
                calls.append(path)
                return "Visible board revision: REV-D4"
            items, manifests = parse_corpus(corpus,config_at(temp),fake_vision,time.monotonic()+30)
            contents = "\n".join(e.content for e in items)
            self.assertIn("Sampling: Q2 FY34",contents)
            self.assertIn("FAN-8741",contents)
            self.assertIn("TIMEOUT = 235",contents)
            self.assertIn("REV-D4",contents)
            self.assertNotIn("never retrieve this",contents)
            self.assertEqual(len(calls),2)
            status = {m["path"]:m["status"] for m in manifests}
            self.assertEqual(status["locked.pdf"],"failed")
            self.assertEqual(status["broken.docx"],"failed")
            self.assertEqual(status["unknown.dat"],"skipped")
            if os.getuid() != 0:
                self.assertEqual(status["denied.txt"],"failed")
            self.assertTrue(any(e.location.get("sheet") == "Second" for e in items))
            self.assertTrue(any(e.location.get("page") == 1 for e in items))
            self.assertTrue(all(len(e.chunk_id)==24 and not e.source_path.startswith("/") for e in items))

    def test_parse_timeout_continues(self):
        with tempfile.TemporaryDirectory() as temp:
            temp=Path(temp)
            corpus=temp/"corpus"
            corpus.mkdir()
            (corpus/"bad.txt").write_text("bad")
            (corpus/"good.txt").write_text("good")
            original=subprocess.run
            def run(command, **kwargs):
                if Path(command[-2]).name == "bad.txt":
                    raise subprocess.TimeoutExpired(command,kwargs["timeout"])
                return original(command,**kwargs)
            with patch("rag.parsers.subprocess.run",side_effect=run):
                evidence,manifest=parse_corpus(corpus,config_at(temp),lambda *_:"",time.monotonic()+5)
            self.assertEqual([e.source_path for e in evidence],["good.txt"])
            self.assertEqual(manifest[0]["status"],"failed")


class ConfigAndSnapshotTests(unittest.TestCase):
    def test_dotenv_is_read_with_environment_precedence_no_execution(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp,"example.env")
            path.write_text('RAG_TEST_A="from file"\nRAG_TEST_B=$(touch /never)\n')
            with patch.dict(os.environ,{"RAG_TEST_A":"exported"}):
                load_env(path)
                self.assertEqual(os.environ["RAG_TEST_A"],"exported")
                self.assertEqual(os.environ["RAG_TEST_B"],"$(touch /never)")

    def test_weight_manifest_rejects_modified_file(self):
        import hashlib
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)
            (path/"weight.safetensors").write_bytes(b"weights")
            data={"role":"generator","repo_id":"test/candidate","revision":"a"*40,
                  "files":{"weight.safetensors":{"size":7,"sha256":hashlib.sha256(b"weights").hexdigest()}}}
            (path/"weights-manifest.json").write_text(json.dumps(data))
            verify_snapshot(path,"generator")
            (path/"weight.safetensors").write_bytes(b"changed")
            with self.assertRaises(ValueError): verify_snapshot(path,"generator")


class DynamicCorpusTests(unittest.TestCase):
    def test_first_index_includes_model_load_in_budget(self):
        with tempfile.TemporaryDirectory() as temp:
            temp=Path(temp)
            corpus=temp/"empty"
            corpus.mkdir()
            pipeline=Pipeline(config_at(temp))
            pipeline.started=time.monotonic()-600
            with self.assertRaises(TimeoutError): pipeline.index(corpus,time.monotonic()+5)
            self.assertIsNone(pipeline.retriever)

    def test_rebuild_switch_and_failed_index_do_not_reuse_evidence(self):
        with tempfile.TemporaryDirectory() as temp:
            temp=Path(temp)
            a,b=temp/"a",temp/"b"
            a.mkdir(); b.mkdir()
            (a/"one.txt").write_text("PX-713 timeout: 231")
            (b/"entirely-different.txt").write_text("PX-918 timeout: 297")
            pipeline=Pipeline(config_at(temp))
            pipeline.index(a,time.monotonic()+5)
            result,_=pipeline.query(a,"What is the PX-713 timeout?",time.monotonic()+5)
            self.assertEqual(result["answer"],"231")
            with self.assertRaises(ValueError): pipeline.query(b,"PX-918 timeout",time.monotonic()+5)
            pipeline.index(b,time.monotonic()+5)
            result,_=pipeline.query(b,"What is the PX-918 timeout?",time.monotonic()+5)
            self.assertEqual(result,{"answer":"297","citations":["entirely-different.txt"]})
            (b/"entirely-different.txt").write_text("PX-918 timeout: 298")
            with self.assertRaises(ValueError): pipeline.query(b,"PX-918 timeout",time.monotonic()+5)
            pipeline.index(b,time.monotonic()+5)
            self.assertEqual(pipeline.query(b,"PX-713 timeout",time.monotonic()+5)[0]["answer"],"")
            with self.assertRaises(FileNotFoundError): pipeline.index(temp/"missing",time.monotonic()+5)
            self.assertIsNone(pipeline.retriever)

    def test_empty_corpus_and_restore(self):
        with tempfile.TemporaryDirectory() as temp:
            temp=Path(temp)
            corpus=temp/"empty"
            corpus.mkdir()
            cfg=config_at(temp)
            pipeline=Pipeline(cfg)
            pipeline.index(corpus,time.monotonic()+5)
            self.assertEqual(pipeline.query(corpus,"unknown",time.monotonic()+5)[0],{"answer":"","citations":[]})
            restarted=Pipeline(cfg)
            self.assertTrue(restarted.restore())
            (corpus/"new.txt").write_text("new corpus")
            self.assertFalse(Pipeline(cfg).restore())


class WorkerContractTests(unittest.TestCase):
    def test_index_waits_for_container_worker_loading_and_checks_config(self):
        with tempfile.TemporaryDirectory() as temp:
            temp=Path(temp)
            corpus=temp/"corpus"
            corpus.mkdir()
            (corpus/"data.txt").write_text("PX-713 timeout: 231")
            env=os.environ.copy()
            env.update(RAG_BACKEND="extractive",RAG_DEVICE="cpu",RAG_REQUIRE_ROCM="false",
                       RAG_CONFIG=str(ROOT/"configs/E0.json"),RAG_SOCKET=str(temp/"boot.sock"),
                       RAG_STATE_DIR=str(temp/"state"),RAG_OUTPUT_DIR=str(temp/"output"),RAG_MODEL_DIR=str(temp/"models"))
            code="import time, worker\noriginal=worker.Pipeline\ndef loading(config):\n time.sleep(.6)\n return original(config)\nworker.Pipeline=loading\nworker.main()"
            worker=subprocess.Popen([sys.executable,"-c",code],cwd=ROOT,env=env,
                                    stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
            try:
                deadline=time.monotonic()+5
                while not Path(str(temp/"boot.sock")+".lock").exists() and time.monotonic()<deadline:
                    time.sleep(.01)
                proc=subprocess.run([sys.executable,str(ROOT/"app.py"),"--index",str(corpus)],
                                    cwd=ROOT,env=env,capture_output=True,text=True,timeout=10)
                self.assertEqual(proc.returncode,0,proc.stderr)
                self.assertFalse((temp/"state/worker.log").exists(),"CLI must wait for CMD worker")
                alternate=json.loads((ROOT/"configs/E0.json").read_text())
                alternate.update(chunk_size=256,chunk_overlap=32)
                (temp/"other.json").write_text(json.dumps(alternate))
                env["RAG_CONFIG"]=str(temp/"other.json")
                proc=subprocess.run([sys.executable,str(ROOT/"app.py"),"--corpus",str(corpus),
                                     "--query-id","mismatch","--query","PX-713 timeout"],cwd=ROOT,env=env,
                                    capture_output=True,text=True,timeout=10)
                self.assertEqual(proc.returncode,1)
                self.assertIn("configuration mismatch",proc.stderr)
                self.assertEqual(json.loads((temp/"output/mismatch_output.json").read_text()),{"answer":"","citations":[]})
            finally:
                worker.terminate()
                worker.wait(timeout=5)

    def test_separate_process_queries_same_worker_and_atomic_outputs(self):
        with tempfile.TemporaryDirectory() as temp:
            temp=Path(temp)
            corpus=temp/"corpus"
            corpus.mkdir()
            (corpus/"renamed.txt").write_text("PX-713 timeout: 231")
            env=os.environ.copy()
            env.update(RAG_BACKEND="extractive", RAG_DEVICE="cpu", RAG_REQUIRE_ROCM="false",
                       RAG_CONFIG=str(ROOT/"configs/E0.json"),RAG_SOCKET=str(temp/"w.sock"),
                       RAG_STATE_DIR=str(temp/"state"),RAG_OUTPUT_DIR=str(temp/"output"),
                       RAG_MODEL_DIR=str(temp/"models"))
            worker=subprocess.Popen([sys.executable,str(ROOT/"worker.py")],cwd=ROOT,env=env,
                                    stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
            try:
                deadline=time.monotonic()+10
                while time.monotonic()<deadline:
                    try:
                        ping=request(str(temp/"w.sock"),{"op":"ping"},.5)
                        break
                    except OSError: time.sleep(.05)
                else: self.fail("worker did not start")
                def cli(*args):
                    return subprocess.run([sys.executable,str(ROOT/"app.py"),*args],cwd=ROOT,env=env,
                                          capture_output=True,text=True,timeout=10)
                self.assertEqual(cli("--index",str(corpus)).returncode,0)
                for qid,query,answer in [("evaluator_63","What is the PX-713 timeout?","231"),
                                         ("evaluator_84","What is the PX-713 price?","")]:
                    proc=cli("--corpus",str(corpus),"--query-id",qid,"--query",query)
                    self.assertEqual(proc.returncode,0,proc.stderr)
                    out=json.loads((temp/"output"/(qid+"_output.json")).read_text())
                    self.assertEqual(out["answer"],answer)
                    self.assertEqual(request(str(temp/"w.sock"),{"op":"ping"},1)["pid"],ping["pid"])
                self.assertEqual(request(str(temp/"w.sock"),{"op":"ping"},1)["generation"],1)
                (corpus/"renamed.txt").write_text("PX-713 timeout: 244")
                proc=cli("--corpus",str(corpus),"--query-id","stale","--query","PX-713 timeout")
                self.assertNotEqual(proc.returncode,0)
                self.assertEqual(json.loads((temp/"output/stale_output.json").read_text()),{"answer":"","citations":[]})
                bad=cli("--corpus",str(corpus),"--query-id","../escape","--query","unknown")
                self.assertNotEqual(bad.returncode,0)
                self.assertIsNone(worker.poll())
            finally:
                worker.terminate()
                worker.wait(timeout=5)

    def test_missing_worker_writes_output_does_not_bootstrap(self):
        with tempfile.TemporaryDirectory() as temp:
            temp=Path(temp)
            env=os.environ.copy()
            env.update(RAG_BACKEND="extractive",RAG_REQUIRE_ROCM="false",RAG_DEVICE="cpu",
                       RAG_SOCKET=str(temp/"none.sock"),RAG_STATE_DIR=str(temp/"state"),RAG_OUTPUT_DIR=str(temp/"output"),
                       RAG_CONFIG=str(ROOT/"configs/E0.json"))
            proc=subprocess.run([sys.executable,str(ROOT/"app.py"),"--corpus",str(temp),
                                 "--query-id","new_qid","--query","unknown"],env=env,capture_output=True,timeout=5)
            self.assertEqual(proc.returncode,1)
            self.assertEqual(json.loads((temp/"output/new_qid_output.json").read_text()),{"answer":"","citations":[]})
            self.assertFalse((temp/"none.sock").exists())


if __name__ == "__main__":
    unittest.main()
