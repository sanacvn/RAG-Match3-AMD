"""Parse untrusted files in a bounded subprocess. Never execute corpus Python."""
import csv
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time

from .schema import Evidence

SUPPORTED = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".log", ".py", ".png", ".jpg", ".jpeg"}


def inventory(root):
    entries = []
    errors = []
    def onerror(exc):
        errors.append({"path": os.path.relpath(exc.filename, root), "status": "failed", "error": type(exc).__name__})
    for folder, dirs, files in os.walk(root, followlinks=False, onerror=onerror):
        dirs[:] = sorted(d for d in dirs if not Path(folder, d).is_symlink())
        for name in sorted(files):
            p = Path(folder, name)
            relative = p.relative_to(root).as_posix()
            try:
                st = p.lstat()
                entries.append((relative, st.st_size, st.st_mtime_ns, st.st_mode, p.is_symlink()))
            except OSError as exc:
                errors.append({"path": relative, "status": "failed", "error": type(exc).__name__})
    return entries, errors


def fingerprint(root):
    entries, errors = inventory(root)
    return hashlib.sha256(json.dumps([entries, errors], sort_keys=True).encode()).hexdigest()


def _table(headers, row):
    return " | ".join(f"{h or 'column_' + str(i+1)}: {v}" for i, (h, v) in enumerate(zip(headers, row)))


def parse_native(path, scratch):
    """Child-only parser. Return ordered units and rendered image jobs."""
    units, images = [], []
    suffix = path.suffix.lower()
    def add(text, location, structured=False):
        if text and text.strip():
            units.append({"content": text.strip(), "location": location, "structured": structured})
    if suffix in {".txt", ".log", ".py"}:
        # Reading .py as text is intentional; imports/exec/eval are forbidden.
        add(path.read_text(errors="replace"), {"line_start": 1})
    elif suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig", errors="replace") as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            rows = csv.reader(f, dialect)
            headers = next(rows, [])
            for n, row in enumerate(rows, 2):
                width = max(len(headers), len(row))
                h = headers + [f"column_{i+1}" for i in range(len(headers), width)]
                add(_table(h, row + [""] * (width-len(row))), {"row": n}, True)
    elif suffix == ".xlsx":
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book:
                headers = None
                for n, raw in enumerate(sheet.iter_rows(values_only=True), 1):
                    row = ["" if v is None else str(v) for v in raw]
                    if not any(row):
                        continue
                    if headers is None:
                        headers = row
                        continue
                    add(f"Sheet: {sheet.title}\n" + _table(headers, row), {"sheet": sheet.title, "row": n}, True)
        finally:
            book.close()
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(path)
        context = []
        for n, node in enumerate(doc.element.body):
            if node.tag.endswith("}p"):
                text = Paragraph(node, doc).text
                add("\n".join(context[-2:] + [text]), {"block": n+1})
                if text:
                    context.append(text)
            elif node.tag.endswith("}tbl"):
                table = Table(node, doc)
                rows = [[c.text for c in r.cells] for r in table.rows]
                if rows:
                    headers = rows[0]
                    for r, row in enumerate(rows[1:], 2):
                        add("\n".join(context[-2:]) + "\n" + _table(headers, row), {"table": n+1, "row": r}, True)
    elif suffix == ".pdf":
        import pymupdf
        with pymupdf.open(path) as doc:
            if doc.needs_pass:
                raise PermissionError("Encrypted PDF: no password supplied")
            for n, page in enumerate(doc, 1):
                text = page.get_text(sort=True)
                add(text, {"page": n})
                # Tables preserve header/value relationships alongside page prose.
                try:
                    for t, table in enumerate(page.find_tables().tables, 1):
                        rows = table.extract()
                        if rows:
                            headers = [str(x or "") for x in rows[0]]
                            for r, row in enumerate(rows[1:], 2):
                                add(_table(headers, [str(x or "") for x in row]), {"page": n, "table": t, "row": r}, True)
                except (AttributeError, ValueError):
                    pass
                # A scan or a page with figures gets vision during indexing.
                if not text.strip() or page.get_images():
                    target = Path(scratch, f"page-{n}.png")
                    page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(target)
                    images.append({"path": str(target), "location": {"page": n, "region": "whole_page"}})
    elif suffix in {".png", ".jpg", ".jpeg"}:
        images.append({"path": str(path), "location": {"region": "whole_image"}})
    return {"units": units, "images": images}


def chunk_text(text, size, overlap):
    # Whitespace-token approximation; experiment reports must name this tokenizer.
    # Split on lines when possible, retaining log/code location and nearby context.
    lines = text.splitlines(keepends=True)
    if not lines:
        return
    i = 0
    while i < len(lines):
        j, words = i, 0
        while j < len(lines) and (words < size or j == i):
            words += len(lines[j].split())
            j += 1
        block = "".join(lines[i:j]).strip()
        if len(block.split()) > size * 2 and j == i + 1:
            tokens = block.split()
            for start in range(0, len(tokens), size-overlap):
                yield " ".join(tokens[start:start+size]), i+1, j
        else:
            yield block, i+1, j
        if j == len(lines):
            break
        back, count = j, 0
        while back > i+1 and count < overlap:
            back -= 1
            count += len(lines[back].split())
        i = max(i+1, back)


def parse_corpus(root, config, vision, deadline):
    evidence, manifest = [], []
    entries, errors = inventory(root)
    manifest.extend(errors)
    for relative, size, _, _, symlink in entries:
        path = root / relative
        record = {"path": relative}
        start = time.monotonic()
        try:
            if time.monotonic() >= deadline:
                raise TimeoutError("Startup budget exhausted")
            if symlink or path.suffix.lower() not in SUPPORTED:
                record.update(status="skipped", error="symlink or unsupported type")
                continue
            if size > config.max_file_mb * 1024**2:
                record.update(status="skipped", error="file size limit")
                continue
            # Open catches actual access errors; hashing also records content identity.
            digest = hashlib.sha256()
            with path.open("rb") as f:
                for block in iter(lambda: f.read(1024*1024), b""):
                    digest.update(block)
            record["sha256"] = digest.hexdigest()
            with tempfile.TemporaryDirectory(prefix="rag-parse-") as scratch:
                proc = subprocess.run([sys.executable, "-m", "rag.parse_job", str(path), scratch],
                                      capture_output=True, text=True,
                                      timeout=min(config.parse_timeout, max(.01, deadline-time.monotonic())))
                if proc.returncode:
                    raise ValueError(proc.stderr.strip()[-500:] or "parser failed")
                parsed = json.loads(proc.stdout)
                units = parsed["units"]
                failures = []
                for job in parsed["images"]:
                    try:
                        content = vision(Path(job["path"]), deadline)
                        if content:
                            units.append({"content": content, "location": job["location"], "structured": True})
                        else:
                            failures.append("image unreadable or vision disabled")
                    except Exception as exc:
                        failures.append(type(exc).__name__ + ": " + str(exc)[:200])
                count = 0
                for unit in units:
                    blocks = [(unit["content"], 1, 1)] if unit["structured"] else chunk_text(unit["content"], config.chunk_size, config.chunk_overlap)
                    for content, line_start, line_end in blocks:
                        location = dict(unit["location"])
                        if "line_start" in location:
                            location.update(line_start=line_start, line_end=line_end)
                        cid = hashlib.sha256(json.dumps([relative, location, content], sort_keys=True).encode()).hexdigest()[:24]
                        evidence.append(Evidence(cid, relative, content, path.suffix.lower(), location))
                        count += 1
                record.update(status="partial" if failures else "parsed", chunks=count)
                if failures:
                    record["errors"] = failures
        except Exception as exc:
            record.update(status="failed", error=type(exc).__name__ + ": " + str(exc)[:500])
        finally:
            record["seconds"] = time.monotonic()-start
            manifest.append(record)
    return evidence, manifest
